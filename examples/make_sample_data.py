#!/usr/bin/env python3
"""
Builds a fictional sample submission so make_packet.py can be tried without real student work.

Writes, next to this script:
    sample_course_data.sav                      the fictional "course" data set (N = 60)
    sample_submission/Sample_Workbook.docx      a filled-in workbook with charts and answers
    sample_submission/Sample_Data.sav           the student's copy of the data, with value labels

Every name and number here is invented. Requires python-docx, pyreadstat, pandas, pillow.
"""
import io
import os
import random

import docx
import pandas as pd
import pyreadstat
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "sample_submission")
random.seed(2026)

N = 60
df = pd.DataFrame({
    "StudyHours": [max(0, round(random.gauss(9, 4))) for _ in range(N)],
    "Housing": [random.choice([1, 1, 2, 2, 2, 3]) for _ in range(N)],
    "CreditHours": [random.choice([6, 9, 12, 12, 12, 15, 15, 18]) for _ in range(N)],
    "Major": [random.choice([1, 2, 2, 3, 4, 4, 5]) for _ in range(N)],
    "CommuteMin": [max(0, round(random.gauss(22, 11))) for _ in range(N)],
})
df = df.astype(float)

labels = {
    "Housing": {1: "On campus", 2: "Off campus with family", 3: "Off campus other"},
    "Major": {1: "Biology", 2: "Business", 3: "Education", 4: "Nursing", 5: "Other"},
}
varlabels = ["Weekly study hours", "Housing type", "Credit hours this term", "Declared major",
             "One-way commute in minutes"]

os.makedirs(OUT, exist_ok=True)
pyreadstat.write_sav(df, os.path.join(HERE, "sample_course_data.sav"))
pyreadstat.write_sav(df, os.path.join(OUT, "Sample_Data.sav"), column_labels=varlabels,
                     variable_value_labels=labels)


def chart_png(title):
    im = Image.new("RGB", (320, 200), "white")
    d = ImageDraw.Draw(im)
    for i, h in enumerate([60, 120, 90, 40]):
        d.rectangle([30 + i * 70, 180 - h, 80 + i * 70, 180], fill=(70, 110, 170))
    d.text((10, 5), title, fill="black")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def box(doc, text=None, image=None):
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    p = cell.paragraphs[0]
    if image is not None:
        p.add_run().add_picture(image, width=docx.shared.Inches(2.5))
    if text:
        cell.add_paragraph(text)
    doc.add_paragraph("")


doc = docx.Document()
doc.add_paragraph("Your Name: Sample Student")
vt = doc.add_table(rows=6, cols=3)
rows = [("Variable", "Quantitative or categorical", "Scale of measurement"),
        ("StudyHours", "Quantitative", "Ratio"), ("Housing", "Categorical", "Nominal"),
        ("CreditHours", "Quantitative", "Ratio"), ("Major", "Categorical", "Nominal"),
        ("CommuteMin", "Quantitative", "Ratio")]
for r, vals in zip(vt.rows, rows):
    for c, v in zip(r.cells, vals):
        c.text = v
doc.add_paragraph("")
box(doc, image=chart_png("Bar chart, StudyHours"))
box(doc, image=chart_png("Histogram, StudyHours"))
box(doc, image=chart_png("Boxplot, StudyHours"))
box(doc, text="The histogram shows the shape best. Study hours center near 9 with a longer right tail. "
              "The bar chart is the weakest display because it treats each hour value as its own category.")
box(doc, text="Most students live off campus with family.", image=chart_png("Housing"))
box(doc, text="Twelve credit hours is the most common load.", image=chart_png("CreditHours"))
box(doc, text="Business and Nursing are the largest majors.", image=chart_png("Major"))
box(doc, text="Commute times center near 20 minutes and are roughly symmetric.", image=chart_png("CommuteMin"))
box(doc, text="Students in this sample typically study about 9 hours a week, carry 12 credit hours, "
              "and commute about 20 minutes each way. Most live off campus with family. "
              "Business and Nursing are the most common majors.")
box(doc, text="I learned to match the display to the scale of measurement and to read skew from a boxplot. "
              "I practiced building charts and adding value labels in SPSS. "
              "At work I could use these displays to summarize survey results. "
              "How would the picture change with a larger sample?")
doc.save(os.path.join(OUT, "Sample_Workbook.docx"))
print("wrote sample data to", HERE)
