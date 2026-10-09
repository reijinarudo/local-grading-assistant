#!/usr/bin/env python3
"""
make_packet.py  --  evidence packet builder for SPSS homework submissions

Turns one student's submission (workbook .docx, data .sav, output .spv) into a single
plain-text evidence packet that a local language model can grade from. No model is
involved here. Everything in the packet is extracted or recomputed; nothing is judged.

Everything specific to one assignment (answer box names, variables, codebook, the true
values of the shared data set, report keywords) lives in a JSON course config file.
See course_config.example.json.

USAGE
    python3 make_packet.py <folder> --config course_config.json
    python3 make_packet.py a.docx b.sav c.spv --config course_config.json
    python3 make_packet.py <folder> --out packets            write packet + chart images under ./packets
    python3 make_packet.py <folder> --name "Student A"       override the student label used in filenames
    python3 make_packet.py <folder> --template blank_workbook.docx --source-sav course_data.sav

    The config path can also come from the GRADER_CONFIG environment variable. Without
    either, the script looks for course_config.json, then course_config.example.json,
    next to itself.

OUTPUT
    <out>/<Name>_packet.txt        the evidence packet (feed this to the model)
    <out>/<Name>_charts/           every chart image, named by the box it was pasted into

Dependencies: python-docx, pyreadstat, and the standard library.
Optional: pillow (TIFF to PNG), openpyxl (Excel submissions), pypdf (PDF workbooks).
If an optional library is absent, those files are listed but not parsed and the packet says so.
"""
import argparse
import datetime as _dt
import hashlib
import json
import math
import os
import re
import sys
import zipfile
from collections import Counter

# --------------------------------------------------------------------------------------
# Assignment facts. Loaded from a JSON course config so no answer key lives in the code.
# --------------------------------------------------------------------------------------

HERE = os.path.dirname(os.path.abspath(__file__))

ASSIGNMENT_TITLE = "Assignment"
BOX_NAMES = {}
FIVE_VARS = []          # the analysis variables, in the order they appear in the packet
CATEGORICAL = set()
CODEBOOK = {}
TRUE = {}
REPORT_KEYWORDS = {}
NAMED_GROUP_KEYWORDS = []
NAMED_GROUP_LABEL = "a specific group"
IGNORE_MISSING_COLUMNS = []
REPORT_WORD_LIMIT = 150
LAYOUT = {
    "variable_table_box": 0,
    "image_only_boxes": [1, 2, 3],
    "comparison_box": 4,
    "chart_commentary_boxes": [5, 6, 7, 8],
    "report_box": 9,
    "reflection_box": 10,
}


def _int_keys(d):
    return {int(k): v for k, v in d.items()}


def find_config(explicit=None):
    for cand in (explicit, os.environ.get("GRADER_CONFIG"),
                 os.path.join(HERE, "course_config.json"),
                 os.path.join(HERE, "course_config.example.json")):
        if cand and os.path.exists(cand):
            return cand
    raise ValueError("No course config found. Pass --config or set GRADER_CONFIG.")


def load_config(path=None):
    """Read the course config and set the module-level assignment facts."""
    global ASSIGNMENT_TITLE, BOX_NAMES, FIVE_VARS, CATEGORICAL, CODEBOOK, TRUE
    global REPORT_KEYWORDS, NAMED_GROUP_KEYWORDS, NAMED_GROUP_LABEL, IGNORE_MISSING_COLUMNS
    global REPORT_WORD_LIMIT, LAYOUT
    path = find_config(path)
    with open(path, encoding="utf8") as fh:
        c = json.load(fh)
    ASSIGNMENT_TITLE = c.get("assignment_title", ASSIGNMENT_TITLE)
    BOX_NAMES = _int_keys(c.get("box_names", {}))
    FIVE_VARS = list(c["variables"])
    CATEGORICAL = set(c.get("categorical", []))
    CODEBOOK = {v: _int_keys(codes) for v, codes in c.get("codebook", {}).items()}
    TRUE = {}
    for v, t in c.get("true_values", {}).items():
        TRUE[v] = _int_keys(t) if v in CATEGORICAL else t
    REPORT_KEYWORDS = c.get("report_keywords", {})
    NAMED_GROUP_KEYWORDS = c.get("named_group_keywords", [])
    NAMED_GROUP_LABEL = c.get("named_group_label", NAMED_GROUP_LABEL)
    IGNORE_MISSING_COLUMNS = c.get("ignore_missing_columns", [])
    REPORT_WORD_LIMIT = int(c.get("report_word_limit", REPORT_WORD_LIMIT))
    LAYOUT = dict(LAYOUT, **c.get("layout", {}))
    return path


# --------------------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------------------

def wc(text):
    return len(text.split())


def fmt(x, nd=2):
    if x is None:
        return "n/a"
    if isinstance(x, float):
        return ("%%.%df" % nd) % x
    return str(x)


def tukey_hinges(values):
    """Quartiles the way SPSS boxplots compute them (Tukey hinges)."""
    s = sorted(values)
    n = len(s)
    if n < 2:
        return (None, None)
    half = (n + 1) // 2  # include the median in both halves when n is odd
    lower = s[:half]
    upper = s[n - half:]

    def med(v):
        m = len(v)
        return v[m // 2] if m % 2 else (v[m // 2 - 1] + v[m // 2]) / 2.0

    return (med(lower), med(upper))


def linear_percentile(values, p):
    """Percentile with linear interpolation (what Excel, pandas, and numpy use by default)."""
    s = sorted(values)
    n = len(s)
    if n == 0:
        return None
    pos = (n - 1) * p
    lo = int(math.floor(pos))
    hi = min(lo + 1, n - 1)
    return s[lo] + (s[hi] - s[lo]) * (pos - lo)


def describe(values):
    v = [float(x) for x in values if x is not None and not (isinstance(x, float) and math.isnan(x))]
    n = len(v)
    if n == 0:
        return None
    mean = sum(v) / n
    s = sorted(v)
    median = s[n // 2] if n % 2 else (s[n // 2 - 1] + s[n // 2]) / 2.0
    cnt = Counter(v)
    top = max(cnt.values())
    modes = sorted(k for k, c in cnt.items() if c == top)
    var = sum((x - mean) ** 2 for x in v) / (n - 1) if n > 1 else 0.0
    sd = math.sqrt(var)
    # SPSS skewness (G1, the adjusted Fisher-Pearson coefficient)
    if n > 2 and sd > 0:
        m2 = sum((x - mean) ** 2 for x in v) / n
        m3 = sum((x - mean) ** 3 for x in v) / n
        g1 = m3 / (m2 ** 1.5)
        G1 = g1 * math.sqrt(n * (n - 1)) / (n - 2)
    else:
        G1 = 0.0
    q1h, q3h = tukey_hinges(v)
    q1l, q3l = linear_percentile(v, .25), linear_percentile(v, .75)
    out = {"n": n, "mean": mean, "median": median, "mode": modes, "mode_n": top, "sd": sd,
           "min": s[0], "max": s[-1], "range": s[-1] - s[0], "skew": G1,
           "q1_hinge": q1h, "q3_hinge": q3h, "q1_lin": q1l, "q3_lin": q3l}
    for tag, q1, q3 in (("hinge", q1h, q3h), ("lin", q1l, q3l)):
        iqr = q3 - q1
        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        out["iqr_" + tag] = iqr
        out["fences_" + tag] = (lo, hi)
        out["outliers_" + tag] = sorted(x for x in v if x < lo or x > hi)
    out["freq"] = sorted(cnt.items())
    return out


def xml_text(tag, xml):
    m = re.search(r"<%s[^>]*>(.*?)</%s>" % (re.escape(tag), re.escape(tag)), xml, re.S)
    return m.group(1).strip() if m else None


# --------------------------------------------------------------------------------------
# Workbook (.docx)
# --------------------------------------------------------------------------------------

def scan_docx_layers(path):
    """Counts of things a PDF render would hide. Nonzero beyond inline images deserves a look."""
    z = zipfile.ZipFile(path)
    names = z.namelist()
    x = z.read("word/document.xml").decode("utf8", "ignore")
    media = sorted(n for n in names if n.startswith("word/media/"))
    # which media files are actually referenced from document.xml
    rels = z.read("word/_rels/document.xml.rels").decode("utf8", "ignore") if "word/_rels/document.xml.rels" in names else ""
    rid_to_target = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    used_rids = set(re.findall(r'r:embed="([^"]+)"', x))
    referenced = {"word/" + rid_to_target[r].lstrip("/").replace("word/", "") for r in used_rids if r in rid_to_target}
    orphans = [m for m in media if m not in referenced]
    return {
        "comments_part": "word/comments.xml" in names,
        "tracked_insertions": x.count("<w:ins "),
        "tracked_deletions": x.count("<w:del "),
        "hidden_text_runs": x.count("<w:vanish"),
        "text_boxes": x.count("txbxContent"),
        "floating_images": x.count("<wp:anchor"),
        "inline_images": x.count("<wp:inline"),
        "cropped_images": x.count("<a:srcRect"),
        "vml_shapes": x.count("<v:shape"),
        "media_files": len(media),
        "orphan_media": orphans,
    }


def docx_forensics(path):
    z = zipfile.ZipFile(path)
    names = z.namelist()
    core = z.read("docProps/core.xml").decode("utf8", "ignore") if "docProps/core.xml" in names else ""
    app = z.read("docProps/app.xml").decode("utf8", "ignore") if "docProps/app.xml" in names else ""
    info = z.getinfo("word/document.xml")
    systems = {0: "MS-DOS/Windows", 3: "Unix", 19: "macOS"}
    f = {
        "creator": xml_text("dc:creator", core),
        "last_modified_by": xml_text("cp:lastModifiedBy", core),
        "revision": xml_text("cp:revision", core),
        "created": xml_text("dcterms:created", core),
        "modified": xml_text("dcterms:modified", core),
        "application": xml_text("Application", app),
        "total_edit_minutes": xml_text("TotalTime", app),
        "pages": xml_text("Pages", app),
        "words_app": xml_text("Words", app),
        "zip_written_on": systems.get(info.create_system, str(info.create_system)),
        "sha256": hashlib.sha256(open(path, "rb").read()).hexdigest()[:16],
    }
    return f


def extract_docx(path, template_paragraphs, chart_dir):
    import docx
    from docx.oxml.ns import qn
    d = docx.Document(path)
    os.makedirs(chart_dir, exist_ok=True)
    for old in os.listdir(chart_dir):  # fresh run, no stale images
        try:
            os.remove(os.path.join(chart_dir, old))
        except OSError:
            pass
    boxes = []
    for i, t in enumerate(d.tables):
        # text, de-duplicated across merged cells
        texts = []
        rows = []
        for r in t.rows:
            row = []
            for c in r.cells:
                ct = c.text.strip()
                row.append(ct)
                if ct and ct not in texts:
                    texts.append(ct)
            rows.append(row)
        text = "\n".join(texts)
        # images, saved in order
        pics = t._tbl.findall(".//" + qn("pic:pic"))
        saved = []
        for k, pic in enumerate(pics):
            blip = pic.find(".//" + qn("a:blip"))
            rid = blip.get(qn("r:embed")) if blip is not None else None
            if rid and rid in d.part.related_parts:
                part = d.part.related_parts[rid]
                ext = os.path.splitext(part.partname)[1] or ".bin"
                fn = "box%02d_img%d%s" % (i, k + 1, ext)
                fp = os.path.join(chart_dir, fn)
                with open(fp, "wb") as fh:
                    fh.write(part.blob)
                if ext.lower() in (".tif", ".tiff", ".emf", ".wmf", ".bmp"):
                    fn = convert_to_png(fp) or fn
                saved.append(fn)
        boxes.append({"index": i, "name": BOX_NAMES.get(i, "Unlabeled box %d" % i), "rows": rows,
                      "text": text, "words": wc(text), "images": saved, "ncols": len(t.columns)})
    # student text that is outside every box (name line, references, stray answers)
    tp = set(template_paragraphs)
    tp_heads = {t[:40] for t in template_paragraphs}
    extra = []
    for p in d.paragraphs:
        t = p.text.strip()
        if t and t not in tp and t[:40] not in tp_heads:
            extra.append(t)
    return boxes, extra


def convert_to_png(fp):
    """TIFF charts (Mac Word) will not display in a browser. Convert if Pillow is available."""
    try:
        from PIL import Image
    except ImportError:
        return None
    try:
        im = Image.open(fp)
        out = os.path.splitext(fp)[0] + ".png"
        im.convert("RGB").save(out)
    except Exception:
        return None
    try:
        os.remove(fp)
    except OSError:
        pass  # some mounted folders forbid deletion; the PNG is still there
    return os.path.basename(out)


def template_paragraphs(template_path):
    if not template_path or not os.path.exists(template_path):
        return []
    import docx
    return [p.text.strip() for p in docx.Document(template_path).paragraphs if p.text.strip()]


# --------------------------------------------------------------------------------------
# Data file (.sav)
# --------------------------------------------------------------------------------------

def sav_dashboard_dates(path):
    raw = open(path, "rb").read().decode("latin1", "ignore")
    return sorted(set(re.findall(r"DashboardInfo\('(\d{4}-\d{2}-\d{2})", raw)))


def extract_sav(path, source_path=None):
    import pyreadstat
    df, meta = pyreadstat.read_sav(path)
    out = {"file": os.path.basename(path), "n_rows": int(df.shape[0]), "n_cols": int(df.shape[1]),
           "columns": list(df.columns)}
    vl = meta.variable_value_labels or {}
    cl = meta.column_names_to_labels or {}
    out["value_labels_present_for"] = sorted(k for k, v in vl.items() if v)
    out["variable_labels_present_for"] = sorted(k for k, v in cl.items() if v)
    out["measure"] = dict(meta.variable_measure or {})
    other = [c for c in df.columns if c not in IGNORE_MISSING_COLUMNS]
    out["missing_cells"] = int(df[other].isna().sum().sum())
    out["dashboard_dates"] = sav_dashboard_dates(path)
    stats = {}
    for v in FIVE_VARS:
        if v not in df.columns:
            stats[v] = None
            continue
        vals = [x for x in df[v].tolist() if x is not None and not (isinstance(x, float) and math.isnan(x))]
        stats[v] = describe(vals)
    out["stats"] = stats
    if source_path and os.path.exists(source_path):
        a = open(path, "rb").read()
        b = open(source_path, "rb").read()
        out["byte_identical_to_source"] = (a == b)
        try:
            sdf, smeta = pyreadstat.read_sav(source_path)
            common = [c for c in FIVE_VARS if c in sdf.columns and c in df.columns]
            same = all(sorted(sdf[c].tolist()) == sorted(df[c].tolist()) for c in common)
            out["values_match_source"] = same
        except Exception as e:  # pragma: no cover
            out["values_match_source"] = "could not compare (%s)" % e
    return out


# --------------------------------------------------------------------------------------
# Output file (.spv)
# --------------------------------------------------------------------------------------

def extract_spv(path):
    z = zipfile.ZipFile(path)
    cmds, titles = [], []
    for n in sorted(x for x in z.namelist() if "heading" in x.lower() and x.lower().endswith(".xml")):
        x = z.read(n).decode("utf8", "ignore")
        c = re.findall(r'commandName="([^"]+)"', x)
        if c:
            cmds.append(c[0])
        # chart / table labels inside the outline
        for lab in re.findall(r'<label>([^<]{2,80})</label>', x):
            lab = lab.strip()
            if lab not in ("Output", "Title", "Notes", "Active Dataset", "Case Processing Summary", "Statistics", "Log"):
                titles.append(lab)
    charts = sum(1 for n in z.namelist() if n.lower().endswith("chartdata.bin") or "chart" in n.lower())
    return {"file": os.path.basename(path), "procedures_in_order": cmds, "procedure_counts": dict(Counter(cmds)),
            "labels": titles, "chart_entries": charts}


# --------------------------------------------------------------------------------------
# Optional: Excel data and PDF workbook (best effort, no hard dependency)
# --------------------------------------------------------------------------------------

def extract_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        return {"file": os.path.basename(path), "note": "openpyxl not installed; Excel file listed but not parsed."}
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {"file": os.path.basename(path), "sheets": wb.sheetnames, "stats": {}}
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return out
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    for v in FIVE_VARS:
        if v in header:
            j = header.index(v)
            vals = [r[j] for r in rows[1:] if r[j] is not None and isinstance(r[j], (int, float))]
            out["stats"][v] = describe(vals)
    return out


def extract_pdf_text(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    r = PdfReader(path)
    return "\n".join((p.extract_text() or "") for p in r.pages)


# --------------------------------------------------------------------------------------
# Automated flags. Plain observations the model would otherwise have to hunt for.
# --------------------------------------------------------------------------------------

def report_flags(boxes, sav):
    flags = []
    by = {b["index"]: b for b in boxes}
    rep = by.get(LAYOUT["report_box"], {}).get("text", "")
    ref = by.get(LAYOUT["reflection_box"], {}).get("text", "")
    low = rep.lower()
    if wc(rep) == 0:
        flags.append("REPORT box is empty.")
    else:
        flags.append("Report word count: %d (limit %d). %s" % (wc(rep), REPORT_WORD_LIMIT, "OVER LIMIT." if wc(rep) > REPORT_WORD_LIMIT else "Within limit."))
        for char, keys in REPORT_KEYWORDS.items():
            hit = any(k in low for k in keys)
            flags.append("Report mentions %s: %s" % (char, "yes (keyword match, verify it reports a finding, not just a mention)" if hit else "NO keyword found"))
        if NAMED_GROUP_KEYWORDS:
            named = [g for g in NAMED_GROUP_KEYWORDS if g in low]
            flags.append("Report names %s: %s" % (NAMED_GROUP_LABEL, (", ".join(named)) if named else "NO. If the report only lists the topic, the finding was not reported."))
        raw_names = [v for v in FIVE_VARS if v.lower() in low]
        if raw_names:
            flags.append("Report uses raw variable names: %s (check that each is defined)." % ", ".join(raw_names))
        first = FIVE_VARS[0] if FIVE_VARS else None
        true_mean = (TRUE.get(first) or {}).get("mean") if first and first not in CATEGORICAL else None
        if true_mean is not None:
            for num in re.findall(r"\d+\.\d+", rep):
                x = float(num)
                if x != round(true_mean, len(num.split(".")[1])) and abs(x - true_mean) <= 0.1:
                    flags.append("Report contains a %s mean near but not equal to %s (%s)." % (first, true_mean, num))
                    break
    if wc(ref) == 0:
        flags.append("REFLECTION box is empty.")
    else:
        flags.append("Reflection word count: %d. Contains a question mark: %s." % (wc(ref), "yes" if "?" in ref else "NO"))
    # images per box
    image_only = LAYOUT["image_only_boxes"]
    chart_commentary = LAYOUT["chart_commentary_boxes"]
    for i in list(image_only) + list(chart_commentary):
        b = by.get(i)
        if b is None:
            continue
        if not b["images"] and i in image_only:
            flags.append("Box %d (%s) has NO image." % (i, b["name"]))
        if not b["images"] and i in chart_commentary:
            flags.append("Box %d (%s) has NO image (student may have pasted text output instead)." % (i, b["name"]))
        if len(b["images"]) > 1:
            flags.append("Box %d (%s) has %d images where one was asked for." % (i, b["name"], len(b["images"])))
        if "Stem-and-Leaf" in b["text"] or re.search(r"Frequency\s+Stem\s*&\s*Leaf", b["text"]):
            flags.append("Box %d (%s) contains pasted SPSS stem-and-leaf TEXT, which is not one of the three displays the workbook names." % (i, b["name"]))
        if i in chart_commentary and b["words"] == 0:
            flags.append("Box %d (%s) has an image but no commentary." % (i, b["name"]))
    cb = LAYOUT["comparison_box"]
    if cb is not None and by.get(cb) and by[cb]["words"] == 0:
        flags.append("Box %d (%s) is empty." % (cb, by[cb]["name"]))
    vb = LAYOUT["variable_table_box"]
    if vb is not None and by.get(vb):
        empties = sum(1 for r in by[vb]["rows"][1:] for c in r[1:] if not c)
        if empties:
            flags.append("Variable table has %d empty cell(s)." % empties)
    # sav based
    if sav:
        if not sav["value_labels_present_for"]:
            flags.append("LABELING: NO value labels in the .sav. Categorical charts will show numeric codes instead of category names.")
        if not sav["variable_labels_present_for"]:
            flags.append("LABELING: NO variable labels in the .sav. Chart titles and axes carry raw variable names.")
        wrong_measure = [v for v in FIVE_VARS if v not in CATEGORICAL and sav["measure"].get(v) == "nominal"]
        if wrong_measure:
            flags.append("SPSS Measure column set to Nominal for quantitative variable(s): %s (forward-looking note)." % ", ".join(wrong_measure))
        if sav.get("byte_identical_to_source") is True:
            flags.append("Data file is byte identical to the course source file (student never saved changes; combine with other signals before drawing any conclusion).")
        if sav.get("values_match_source") is False:
            flags.append("DATA VALUES DIFFER from the course source file. Recomputed statistics below are from the student's file; treat any disagreement with the true values as a finding.")
    return flags


# --------------------------------------------------------------------------------------
# Packet writer
# --------------------------------------------------------------------------------------

def stat_block(name, st, categorical):
    lines = []
    if st is None:
        return ["  %s: NOT IN FILE" % name]
    if categorical:
        total = st["n"]
        lines.append("  %s (N = %d)" % (name, total))
        truth = TRUE.get(name, {})
        for code, cnt in st["freq"]:
            code_i = int(code)
            lab = CODEBOOK.get(name, {}).get(code_i, "")
            t = truth.get(code_i)
            mark = "" if t is None else ("  MATCH" if t == cnt else "  DIFFERS (true %d)" % t)
            lines.append("    code %d %-34s %3d  (%.1f%%)%s" % (code_i, lab, cnt, 100.0 * cnt / total, mark))
        missing_codes = [c for c in CODEBOOK.get(name, {}) if c not in [int(k) for k, _ in st["freq"]]]
        if missing_codes:
            lines.append("    codes with zero cases (will not appear on an SPSS chart): %s" % missing_codes)
        return lines
    truth = TRUE.get(name, {})

    def chk(key, val, nd=2, tol=0.01):
        t = truth.get(key)
        if t is None:
            return ""
        return "  MATCH" if abs(float(val) - float(t)) <= max(tol, abs(t) * 0.002) else "  DIFFERS (true %s)" % t

    lines.append("  %s (N = %d)" % (name, st["n"]))
    lines.append("    mean %.2f%s" % (st["mean"], chk("mean", st["mean"])))
    lines.append("    median %s%s" % (fmt(st["median"], 1), chk("median", st["median"])))
    if st["mode_n"] > 1 and len(st["mode"]) <= 3:
        lines.append("    mode %s (n = %d)" % (", ".join(fmt(m, 0) for m in st["mode"]), st["mode_n"]))
    else:
        lines.append("    mode: no single repeated value (every value distinct or many ties)")
    lines.append("    SD %.3f%s" % (st["sd"], chk("sd", st["sd"], tol=0.005)))
    lines.append("    min %s  max %s  range %s" % (fmt(st["min"], 0), fmt(st["max"], 0), fmt(st["range"], 0)))
    lines.append("    skewness %.2f (SPSS formula)%s" % (st["skew"], chk("skew", st["skew"], tol=0.02)))
    lines.append("    quartiles, SPSS boxplot method (Tukey hinges): Q1 %s  Q3 %s  IQR %s  fences %s to %s  outliers %s"
                 % (fmt(st["q1_hinge"], 2), fmt(st["q3_hinge"], 2), fmt(st["iqr_hinge"], 2),
                    fmt(st["fences_hinge"][0], 2), fmt(st["fences_hinge"][1], 2),
                    [int(x) if float(x).is_integer() else x for x in st["outliers_hinge"]] or "none"))
    lines.append("    quartiles, Excel/pandas method (linear): Q1 %s  Q3 %s  outliers %s"
                 % (fmt(st["q1_lin"], 2), fmt(st["q3_lin"], 2),
                    [int(x) if float(x).is_integer() else x for x in st["outliers_lin"]] or "none"))
    if len(st["freq"]) <= 25:
        lines.append("    frequencies: " + ", ".join("%s (%d)" % (fmt(v, 0), c) for v, c in st["freq"]))
    return lines


def write_packet(name, files, boxes, extra_paras, layers, forensics, sav, spv, xlsx, pdf_text, flags, out_path):
    L = []
    P = L.append
    P("EVIDENCE PACKET  |  %s  |  Student: %s" % (ASSIGNMENT_TITLE, name))
    P("Generated %s by make_packet.py. Everything below was extracted or recomputed by script, not by a model." % _dt.datetime.now().strftime("%Y-%m-%d %H:%M"))
    P("")
    P("=" * 78)
    P("SECTION 1. FILES RECEIVED")
    P("=" * 78)
    for kind in ("workbook", "data", "output"):
        P("  %-9s %s" % (kind + ":", files.get(kind) or "MISSING"))
    missing = [k for k in ("workbook", "data", "output") if not files.get(k)]
    P("  File Submission (all or nothing): %s" % ("all three present" if not missing else "INCOMPLETE, missing " + ", ".join(missing)))
    P("")
    P("=" * 78)
    P("SECTION 2. AUTOMATED FLAGS (observations to verify against the text below)")
    P("=" * 78)
    for f in flags:
        P("  - " + f)
    if not flags:
        P("  (none)")
    P("")
    P("=" * 78)
    P("SECTION 3. THE STUDENT'S ANSWERS, BOX BY BOX (workbook tables in order)")
    P("=" * 78)
    if extra_paras:
        P("Student text found outside the answer boxes:")
        for e in extra_paras:
            P("  | " + e)
        P("")
    for b in boxes:
        P("-" * 78)
        P("BOX %d: %s" % (b["index"], b["name"]))
        P("  images: %d%s   words: %d" % (len(b["images"]), (" [" + ", ".join(b["images"]) + "]") if b["images"] else "", b["words"]))
        if b["index"] == LAYOUT["variable_table_box"] and b["ncols"] >= 3:
            for r in b["rows"]:
                P("  | " + " || ".join(c.replace("\n", " ") for c in r))
        elif b["text"]:
            for line in b["text"].splitlines():
                P("  | " + line)
        else:
            P("  | [NO TEXT IN THIS BOX]")
    if pdf_text:
        P("-" * 78)
        P("WORKBOOK WAS A PDF. Raw page text follows; box boundaries are not recoverable from a PDF.")
        for line in pdf_text.splitlines():
            if line.strip():
                P("  | " + line.strip())
    P("")
    P("=" * 78)
    P("SECTION 4. DATA FILE CHECK")
    P("=" * 78)
    if sav:
        ign = (" (ignoring %s, configured as blank by design)" % ", ".join(IGNORE_MISSING_COLUMNS)) if IGNORE_MISSING_COLUMNS else ""
        P("  file: %s   rows: %d   columns: %d   empty cells: %d%s" % (sav["file"], sav["n_rows"], sav["n_cols"], sav["missing_cells"], ign))
        P("  columns: %s" % ", ".join(sav["columns"]))
        P("  value labels present for: %s" % (", ".join(sav["value_labels_present_for"]) or "NONE"))
        P("  variable labels present for: %s" % (", ".join(sav["variable_labels_present_for"]) or "NONE"))
        P("  SPSS Measure column: %s" % ", ".join("%s=%s" % (k, v) for k, v in sav["measure"].items() if k in FIVE_VARS))
        if sav.get("dashboard_dates"):
            P("  SPSS work dates recorded inside the file: %s" % ", ".join(sav["dashboard_dates"]))
        if "byte_identical_to_source" in sav:
            P("  byte identical to course source file: %s   same data values as source: %s" % (sav["byte_identical_to_source"], sav.get("values_match_source")))
        P("")
        P("  RECOMPUTED STATISTICS FROM THE STUDENT'S OWN FILE (MATCH / DIFFERS is against the course true values)")
        for v in FIVE_VARS:
            L.extend(stat_block(v, sav["stats"].get(v), v in CATEGORICAL))
    elif xlsx:
        P("  Excel data file: %s" % xlsx["file"])
        if "note" in xlsx:
            P("  " + xlsx["note"])
        else:
            P("  sheets: %s" % ", ".join(xlsx["sheets"]))
            P("  (value labels, variable labels, and Measure do not apply to Excel)")
            for v in FIVE_VARS:
                L.extend(stat_block(v, xlsx["stats"].get(v), v in CATEGORICAL))
    else:
        P("  NO DATA FILE. Nothing can be recomputed. Every number the student reports is unverified.")
    P("")
    P("=" * 78)
    P("SECTION 5. OUTPUT FILE CHECK")
    P("=" * 78)
    if spv:
        P("  file: %s" % spv["file"])
        P("  procedures run, in order: %s" % (", ".join(spv["procedures_in_order"]) or "none found"))
        P("  procedure counts: %s" % json.dumps(spv["procedure_counts"]))
        if spv["labels"]:
            P("  output outline (variable and chart names, in order): %s" % "; ".join(spv["labels"]))
        P("  Note: SPSS charts inside the .spv are binary and cannot be read. Chart appearance is judged from the images the student pasted into the workbook.")
    else:
        P("  NO OUTPUT FILE.")
    P("")
    P("=" * 78)
    P("SECTION 6. WORKBOOK FILE FORENSICS (for the instructor; never mention these in student feedback)")
    P("=" * 78)
    if forensics:
        for k in ("creator", "last_modified_by", "revision", "created", "modified", "application",
                  "total_edit_minutes", "pages", "zip_written_on", "sha256"):
            P("  %-20s %s" % (k + ":", forensics.get(k)))
    if layers:
        P("  hidden layer scan: comments part=%s, tracked ins/del=%d/%d, hidden text runs=%d, text boxes=%d, floating images=%d, inline images=%d, cropped images=%d, VML shapes=%d"
          % (layers["comments_part"], layers["tracked_insertions"], layers["tracked_deletions"], layers["hidden_text_runs"],
             layers["text_boxes"], layers["floating_images"], layers["inline_images"], layers["cropped_images"], layers["vml_shapes"]))
        P("  media files in document: %d   orphaned (present but not displayed): %s" % (layers["media_files"], layers["orphan_media"] or "none"))
        if layers["text_boxes"]:
            P("  WARNING: text boxes present. Their contents are NOT in Section 3. The instructor must open the file.")
    P("")
    P("END OF PACKET")
    with open(out_path, "w", encoding="utf8") as fh:
        fh.write("\n".join(L) + "\n")
    return out_path


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------

def collect(paths):
    files = {}
    for p in paths:
        if os.path.isdir(p):
            for fn in sorted(os.listdir(p)):
                collect_one(os.path.join(p, fn), files)
        else:
            collect_one(p, files)
    return files


def collect_one(p, files):
    ext = os.path.splitext(p)[1].lower()
    base = os.path.basename(p)
    if base.startswith("~$") or base.startswith("."):
        return
    if ext == ".docx":
        files.setdefault("workbook", p)
    elif ext == ".pdf":
        files.setdefault("workbook", p)
    elif ext == ".sav":
        files.setdefault("data", p)
    elif ext == ".xlsx":
        files.setdefault("data", p)
    elif ext == ".spv":
        files.setdefault("output", p)


def guess_name(files, boxes, extra):
    for e in extra or []:
        m = re.match(r"Your Name:\s*(.+)", e)
        if m:
            return m.group(1).strip()
    for k in ("workbook", "data", "output"):
        if files.get(k):
            return os.path.splitext(os.path.basename(files[k]))[0].split("_")[0].split(".")[0]
    return "student"


def build_packet(paths, out_dir, name=None, template=None, source_sav=None, config=None):
    """Library entry point. Returns (packet_text, packet_path, chart_dir_or_None).
    Used by the Open WebUI filter; main() below wraps it for the command line."""
    load_config(config)
    files = collect(paths if isinstance(paths, (list, tuple)) else [paths])
    if not files:
        raise ValueError("No .docx, .pdf, .sav, .xlsx, or .spv files found.")
    os.makedirs(out_dir, exist_ok=True)

    boxes, extra, layers, forensics, pdf_text = [], [], None, None, None
    wb = files.get("workbook")
    tmp_name = name or "student"
    chart_dir = os.path.join(out_dir, tmp_name + "_charts")
    if wb and wb.lower().endswith(".docx"):
        tp = template_paragraphs(template)
        boxes, extra = extract_docx(wb, tp, chart_dir)
        layers = scan_docx_layers(wb)
        forensics = docx_forensics(wb)
    elif wb and wb.lower().endswith(".pdf"):
        pdf_text = extract_pdf_text(wb)
        if pdf_text is None:
            pdf_text = "(pypdf not installed; PDF workbook not parsed)"

    name = name or guess_name(files, boxes, extra)
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", name).strip("_") or "student"
    if safe != tmp_name and os.path.isdir(chart_dir):
        new_dir = os.path.join(out_dir, safe + "_charts")
        if not os.path.exists(new_dir):
            os.rename(chart_dir, new_dir)
        chart_dir = new_dir

    sav = xlsx = spv = None
    d = files.get("data")
    if d and d.lower().endswith(".sav"):
        sav = extract_sav(d, source_sav)
    elif d and d.lower().endswith(".xlsx"):
        xlsx = extract_xlsx(d)
    if files.get("output"):
        spv = extract_spv(files["output"])

    flags = report_flags(boxes, sav) if boxes else ["Workbook not parsed box by box (PDF or missing); flags unavailable."]
    out_path = os.path.join(out_dir, safe + "_packet.txt")
    write_packet(name, files, boxes, extra, layers, forensics, sav, spv, xlsx, pdf_text, flags, out_path)
    with open(out_path, encoding="utf8") as fh:
        text = fh.read()
    return text, out_path, (chart_dir if os.path.isdir(chart_dir) else None)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("paths", nargs="+", help="a folder holding one student's files, or the files themselves")
    ap.add_argument("--out", default="./packets")
    ap.add_argument("--name", default=None, help="student label for output filenames")
    ap.add_argument("--template", default=None, help="blank assignment workbook .docx, used to separate student text from instructions")
    ap.add_argument("--source-sav", default=None, help="course data file, to detect an unchanged or altered .sav")
    ap.add_argument("--config", default=None, help="course config JSON (default: $GRADER_CONFIG, then course_config.json next to this script)")
    a = ap.parse_args()
    try:
        text, out_path, chart_dir = build_packet(a.paths, a.out, a.name, a.template, a.source_sav, a.config)
    except ValueError as e:
        sys.exit(str(e))
    print("wrote", out_path)
    if chart_dir:
        print("charts in", chart_dir, "(%d files)" % len(os.listdir(chart_dir)))
    n = len(text.split())
    print("packet length: about %d words (roughly %d tokens)" % (n, int(n * 1.4)))


if __name__ == "__main__":
    main()
